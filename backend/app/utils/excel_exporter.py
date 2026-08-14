import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session, joinedload
from datetime import date
from calendar import monthrange
from io import BytesIO

from app.models.payroll import Payslip, PayrollRun
from app.models.employee import Employee
from app.models.attendance import AttendanceLog, AttendanceStatus
from app.models.department import Department

# ── Brand Colors ─────────────────────────────────────────
BLUE       = "1B4F8E"
LIGHT_BLUE = "D6E4F7"
WHITE      = "FFFFFF"
GRAY       = "F2F4F7"
DARK       = "1A1A2E"
GREEN      = "1E8449"
RED        = "C0392B"


def _header_style(cell, bg=BLUE, font_color=WHITE, bold=True, size=11):
    cell.font      = Font(bold=bold, color=font_color, size=size, name="Arial")
    cell.fill      = PatternFill("solid", fgColor=bg)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _border_all(cell):
    thin = Side(style="thin", color="CCCCCC")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def _set_col_widths(ws, widths: dict):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


# ── 1. PAYROLL REGISTER ──────────────────────────────────
def generate_payroll_register(db: Session, run_id: int) -> BytesIO:
    run = db.query(PayrollRun).filter(PayrollRun.id == run_id).first()
    payslips = db.query(Payslip).options(
        joinedload(Payslip.employee).joinedload(Employee.department)
    ).filter(Payslip.payroll_run_id == run_id).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Payroll Register"
    ws.sheet_view.showGridLines = False

    # Title
    ws.merge_cells("A1:R1")
    title_cell = ws["A1"]
    title_cell.value = f"PAYROLL REGISTER — {run.month}/{run.year}"
    title_cell.font      = Font(bold=True, size=16, color=WHITE, name="Arial")
    title_cell.fill      = PatternFill("solid", fgColor=BLUE)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    ws.merge_cells("A2:R2")
    sub = ws["A2"]
    sub.value     = f"Generated: {date.today()}  |  Status: {run.status.upper()}"
    sub.font      = Font(size=10, color=DARK, name="Arial")
    sub.fill      = PatternFill("solid", fgColor=LIGHT_BLUE)
    sub.alignment = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 20

    # Headers
    headers = [
        "Emp No", "Full Name", "Department",
        "Work Days", "Present", "Absent", "No-Pay Days", "OT Hrs",
        "Basic (LKR)", "OT Amt", "Attend Allow", "Transport", "Meal Allow",
        "Gross Salary", "EPF (8%)", "No-Pay Ded", "Total Deduct", "Net Salary"
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=h)
        _header_style(cell)
        _border_all(cell)
    ws.row_dimensions[3].height = 30

    # Data rows
    totals = {h: 0 for h in range(9, 19)}
    for row_idx, p in enumerate(payslips, 4):
        shade = GRAY if row_idx % 2 == 0 else WHITE
        values = [
            p.employee.employee_no,
            p.employee.full_name,
            p.employee.department.name if p.employee.department else "N/A",
            p.working_days,
            float(p.present_days),
            float(p.absent_days),
            float(p.no_pay_days),
            float(p.ot_hours),
            float(p.basic_salary),
            float(p.ot_amount),
            float(p.attendance_allow),
            float(p.transport_allow),
            float(p.meal_allow),
            float(p.gross_salary),
            float(p.epf_employee),
            float(p.no_pay_deduct),
            float(p.total_deductions),
            float(p.net_salary),
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.fill      = PatternFill("solid", fgColor=shade)
            cell.font      = Font(size=10, name="Arial")
            cell.alignment = Alignment(horizontal="right" if col > 3 else "left")
            _border_all(cell)
            if col >= 9:
                cell.number_format = '#,##0.00'
                totals[col] = totals.get(col, 0) + (val if isinstance(val, float) else 0)

    # Totals row
    total_row = len(payslips) + 4
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True, name="Arial")
    ws.cell(row=total_row, column=1).fill = PatternFill("solid", fgColor=BLUE)
    ws.cell(row=total_row, column=1).font = Font(bold=True, color=WHITE, name="Arial")

    for col in range(2, 9):
        cell = ws.cell(row=total_row, column=col, value="")
        cell.fill = PatternFill("solid", fgColor=BLUE)

    for col in range(9, 19):
        cell = ws.cell(row=total_row, column=col, value=totals.get(col, 0))
        cell.font         = Font(bold=True, color=WHITE, name="Arial")
        cell.fill         = PatternFill("solid", fgColor=BLUE)
        cell.number_format = '#,##0.00'
        cell.alignment    = Alignment(horizontal="right")
        _border_all(cell)

    # Column widths
    _set_col_widths(ws, {
        "A": 10, "B": 22, "C": 16, "D": 8, "E": 8, "F": 8,
        "G": 10, "H": 8,  "I": 14, "J": 12, "K": 12, "L": 12,
        "M": 10, "N": 14, "O": 12, "P": 12, "Q": 13, "R": 14
    })

    # Freeze header rows
    ws.freeze_panes = "A4"

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# ── 2. EPF/ETF REPORT ────────────────────────────────────
def generate_epf_etf_excel(db: Session, run_id: int) -> BytesIO:
    run = db.query(PayrollRun).filter(PayrollRun.id == run_id).first()
    payslips = db.query(Payslip).options(
        joinedload(Payslip.employee)
    ).filter(Payslip.payroll_run_id == run_id).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EPF-ETF Report"
    ws.sheet_view.showGridLines = False

    # Title
    ws.merge_cells("A1:I1")
    t = ws["A1"]
    t.value     = f"EPF / ETF CONTRIBUTION REPORT — {run.month}/{run.year}"
    t.font      = Font(bold=True, size=14, color=WHITE, name="Arial")
    t.fill      = PatternFill("solid", fgColor=BLUE)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    headers = [
        "Emp No", "Full Name", "NIC", "EPF No",
        "Basic Salary", "EPF Employee (8%)",
        "EPF Employer (12%)", "ETF (3%)", "Total EPF"
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=h)
        _header_style(cell)
        _border_all(cell)
    ws.row_dimensions[2].height = 25

    totals = [0] * 9
    for row_idx, p in enumerate(payslips, 3):
        shade = GRAY if row_idx % 2 == 0 else WHITE
        values = [
            p.employee.employee_no,
            p.employee.full_name,
            p.employee.nic,
            p.employee.epf_no or "N/A",
            float(p.basic_salary),
            float(p.epf_employee),
            float(p.epf_employer),
            float(p.etf_employer),
            float(p.epf_employee + p.epf_employer),
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.fill      = PatternFill("solid", fgColor=shade)
            cell.font      = Font(size=10, name="Arial")
            cell.alignment = Alignment(horizontal="right" if col > 4 else "left")
            _border_all(cell)
            if col >= 5:
                cell.number_format = '#,##0.00'
                totals[col - 1] += val if isinstance(val, (int, float)) else 0

    # Totals
    tr = len(payslips) + 3
    for col in range(1, 10):
        cell = ws.cell(row=tr, column=col)
        if col == 1:
            cell.value = "TOTAL"
        elif col >= 5:
            cell.value         = totals[col - 1]
            cell.number_format = '#,##0.00'
        cell.font      = Font(bold=True, color=WHITE, name="Arial")
        cell.fill      = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="right" if col >= 5 else "left")
        _border_all(cell)

    _set_col_widths(ws, {
        "A": 10, "B": 22, "C": 14, "D": 12,
        "E": 14, "F": 16, "G": 16, "H": 12, "I": 14
    })
    ws.freeze_panes = "A3"

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# ── 3. ATTENDANCE REPORT ─────────────────────────────────
def generate_attendance_excel(db: Session, year: int, month: int) -> BytesIO:
    _, last_day = monthrange(year, month)
    start = date(year, month, 1)
    end   = date(year, month, last_day)

    employees = db.query(Employee).filter(Employee.is_active == True).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Report"
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:H1")
    t = ws["A1"]
    t.value     = f"MONTHLY ATTENDANCE REPORT — {month}/{year}"
    t.font      = Font(bold=True, size=14, color=WHITE, name="Arial")
    t.fill      = PatternFill("solid", fgColor=BLUE)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    headers = [
        "Emp No", "Full Name", "Department",
        "Present", "Absent", "On Leave",
        "OT Hours", "Attendance %"
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=h)
        _header_style(cell)
        _border_all(cell)

    for row_idx, emp in enumerate(employees, 3):
        records = db.query(AttendanceLog).filter(
            AttendanceLog.employee_id == emp.id,
            AttendanceLog.date >= start,
            AttendanceLog.date <= end
        ).all()

        present  = sum(1 for r in records if r.status == AttendanceStatus.PRESENT)
        absent   = sum(1 for r in records if r.status == AttendanceStatus.ABSENT)
        on_leave = sum(1 for r in records if r.status == AttendanceStatus.ON_LEAVE)
        ot_hours = float(sum(r.ot_hours or 0 for r in records))
        att_rate = round(present / last_day * 100, 1)

        shade = GRAY if row_idx % 2 == 0 else WHITE
        values = [
            emp.employee_no,
            emp.full_name,
            emp.department.name if emp.department else "N/A",
            present, absent, on_leave,
            ot_hours, att_rate
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.fill      = PatternFill("solid", fgColor=shade)
            cell.font      = Font(size=10, name="Arial")
            cell.alignment = Alignment(horizontal="right" if col > 3 else "left")
            _border_all(cell)
            if col == 8:
                cell.number_format = '0.0"%"'
                # Color code attendance rate
                if att_rate < 85:
                    cell.font = Font(size=10, color=RED, bold=True, name="Arial")
                elif att_rate >= 95:
                    cell.font = Font(size=10, color=GREEN, bold=True, name="Arial")

    _set_col_widths(ws, {
        "A": 10, "B": 22, "C": 16,
        "D": 9, "E": 9, "F": 9,
        "G": 10, "H": 13
    })
    ws.freeze_panes = "A3"

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output